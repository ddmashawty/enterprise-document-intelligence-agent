from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from doc_agent.ingest.loaders import iter_source_files, load_docx, load_file


def _page_break(paragraph) -> None:
    run = paragraph.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def test_load_docx_paragraphs_and_table(tmp_path: Path):
    path = tmp_path / "手册.docx"
    doc = Document()
    doc.add_heading("演示产品参数手册", level=1)
    doc.add_paragraph("TopK：5")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "参数"
    table.rows[0].cells[1].text = "值"
    table.rows[1].cells[0].text = "切片大小"
    table.rows[1].cells[1].text = "500"
    doc.save(path)

    loaded = load_docx(path)
    assert len(loaded.pages) == 1
    assert loaded.pages[0].page == 1
    assert "TopK：5" in loaded.text
    assert "切片大小 | 500" in loaded.text


def test_load_docx_splits_on_page_break(tmp_path: Path):
    path = tmp_path / "分页.docx"
    doc = Document()
    first = doc.add_paragraph("第一页条款")
    _page_break(first)
    doc.add_paragraph("第二页保密等级")
    doc.save(path)

    loaded = load_file(path)
    assert [p.page for p in loaded.pages] == [1, 2]
    assert "第一页条款" in loaded.pages[0].text
    assert "第二页保密等级" in loaded.pages[1].text


def test_load_file_rejects_legacy_doc(tmp_path: Path):
    path = tmp_path / "旧格式.doc"
    path.write_bytes(b"not a real doc")
    with pytest.raises(ValueError, match=r"\.docx"):
        load_file(path)


def test_iter_source_files_includes_docx(tmp_path: Path):
    (tmp_path / "a.DOCX").write_bytes(b"")
    (tmp_path / "b.txt").write_text("x", encoding="utf-8")
    (tmp_path / "skip.csv").write_text("x", encoding="utf-8")
    names = {p.name for p in iter_source_files(tmp_path)}
    assert names == {"a.DOCX", "b.txt"}
