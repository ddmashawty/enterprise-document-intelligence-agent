from __future__ import annotations

import base64
import binascii
import hashlib
import logging
import re
import warnings
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from doc_agent.ingest import tables as tbl
from doc_agent.ingest.ocr import NoOCR, OCRBackend

# pypdf may spam fontTools warnings on complex PDFs
logging.getLogger("pypdf").setLevel(logging.ERROR)
warnings.filterwarnings("ignore", message=".*fontTools.*")


@dataclass
class DocumentPage:
    source: str
    page: int
    text: str


@dataclass
class LoadedDocument:
    source: str
    pages: list[DocumentPage]

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())


@dataclass
class ImageRef:
    source: str
    index: int
    mime: str = ""
    sha256: str = ""
    url: str = ""
    alt: str = ""
    page: int = 1
    data: bytes | None = field(default=None, repr=False)
    needs_ocr: bool = False


@dataclass
class ParsedDocument(LoadedDocument):
    """LoadedDocument plus tables / images / meta.

    ``meta["prose"]`` maps page → text with ``[[TABLE:n]]`` markers so pages can be
    re-rendered after tables change (e.g. redaction); ``meta["ocr_pages"]`` lists
    pages that have no text layer but contain images.
    """

    tables: list[tbl.Table] = field(default_factory=list)
    images: list[ImageRef] = field(default_factory=list)
    meta: dict[str, Any] = field(default_factory=dict)

    @property
    def ocr_pages(self) -> list[int]:
        return list(self.meta.get("ocr_pages") or [])

    @property
    def needs_ocr(self) -> bool:
        return bool(self.ocr_pages) or any(img.needs_ocr for img in self.images)


def as_parsed(doc: LoadedDocument, **meta: Any) -> ParsedDocument:
    if isinstance(doc, ParsedDocument):
        doc.meta.update(meta)
        return doc
    return ParsedDocument(source=doc.source, pages=doc.pages, meta=dict(meta))


def _clean(text: str) -> str:
    lines = [ln.strip() for ln in text.replace("\x00", "").splitlines()]
    lines = [ln for ln in lines if ln]
    return "\n".join(lines)


def load_txt(path: Path) -> LoadedDocument:
    text = _clean(path.read_text(encoding="utf-8", errors="ignore"))
    return LoadedDocument(
        source=str(path),
        pages=[DocumentPage(source=str(path), page=1, text=text)],
    )


def _page_has_image(page: Any) -> bool:
    def walk(resources: Any, depth: int) -> bool:
        if resources is None or depth > 2:
            return False
        xobjects = resources.get_object().get("/XObject")
        if not xobjects:
            return False
        xobjects = xobjects.get_object()
        for name in xobjects:
            obj = xobjects[name].get_object()
            subtype = obj.get("/Subtype")
            if subtype == "/Image":
                return True
            if subtype == "/Form" and walk(obj.get("/Resources"), depth + 1):
                return True
        return False

    try:
        return walk(page.get("/Resources"), 0)
    except Exception:  # noqa: BLE001
        return False


def load_pdf(path: Path, *, tables: bool = False) -> ParsedDocument:
    """pypdf text per page (unchanged behaviour); ``tables=True`` adds pdfplumber tables.

    Pages without a text layer that carry images are listed in ``meta["ocr_pages"]``;
    only a PDF with neither text nor images raises.
    """
    reader = PdfReader(str(path))
    pages: list[DocumentPage] = []
    ocr_pages: list[int] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            raw = ""
        text = _clean(raw)
        if text:
            pages.append(DocumentPage(source=str(path), page=i, text=text))
        elif _page_has_image(page):
            ocr_pages.append(i)
    if not pages and not ocr_pages:
        raise ValueError(f"No extractable text in PDF: {path}")
    doc = ParsedDocument(
        source=str(path),
        pages=pages,
        meta={"format": "pdf", "page_count": len(reader.pages), "ocr_pages": ocr_pages},
    )
    if tables and pages:
        _add_pdf_tables(path, doc)
    return doc


def _pdf_table_grid(table: Any) -> tuple[list[list[str]], list[list[tbl.Origin]]]:
    """pdfplumber reports merged areas as None; copy the covering cell's text by bbox."""
    data = table.extract()
    rows = table.rows
    boxes = [
        (r, c, bbox)
        for r, row in enumerate(rows)
        for c, bbox in enumerate(row.cells)
        if bbox is not None
    ]

    def center_x(c: int) -> float | None:
        xs = [b for _, cc, b in boxes if cc == c]
        return (min(b[0] for b in xs) + max(b[2] for b in xs)) / 2 if xs else None

    def center_y(r: int) -> float | None:
        ys = [b for rr, _, b in boxes if rr == r]
        return (min(b[1] for b in ys) + max(b[3] for b in ys)) / 2 if ys else None

    cells: list[list[str]] = []
    origins: list[list[tbl.Origin]] = []
    for r, row in enumerate(data):
        out_row: list[str] = []
        out_o: list[tbl.Origin] = []
        for c, value in enumerate(row):
            if value is None and rows[r].cells[c] is None:
                cx, cy = center_x(c), center_y(r)
                owner = None
                if cx is not None and cy is not None:
                    owner = next(
                        ((rr, cc) for rr, cc, b in boxes if b[0] <= cx <= b[2] and b[1] <= cy <= b[3]),
                        None,
                    )
                if owner is not None:
                    out_row.append(tbl.join_wrapped(data[owner[0]][owner[1]]))
                    out_o.append(owner)
                    continue
            out_row.append(tbl.join_wrapped(value))
            out_o.append((r, c))
        cells.append(out_row)
        origins.append(out_o)
    return cells, origins


def _add_pdf_tables(path: Path, doc: ParsedDocument) -> None:
    import pdfplumber

    texts = {p.page: p.text for p in doc.pages}
    prose: dict[int, str] = {}
    index = 0
    with pdfplumber.open(str(path)) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            if page_no not in texts:
                continue
            try:
                found = page.find_tables()
            except Exception:  # noqa: BLE001
                found = []
            if not found:
                prose[page_no] = texts[page_no]
                continue
            bboxes = [t.bbox for t in found]

            def outside(obj: dict[str, Any], bboxes: list[Any] = bboxes) -> bool:
                x = (obj.get("x0", 0) + obj.get("x1", 0)) / 2
                y = (obj.get("top", 0) + obj.get("bottom", 0)) / 2
                return not any(b[0] <= x <= b[2] and b[1] <= y <= b[3] for b in bboxes)

            page_tables: list[tbl.Table] = []
            for t in found:
                cells, origins = _pdf_table_grid(t)
                page_tables.append(
                    tbl.Table(cells, source=str(path), page=page_no, index=index, bbox=tuple(t.bbox), origins=origins)
                )
                index += 1
            outside_text = _clean(page.filter(outside).extract_text() or "")
            markers = "\n".join(tbl.TABLE_MARKER.format(index=t.index) for t in page_tables)
            prose[page_no] = f"{outside_text}\n{markers}"
            texts[page_no] = tbl.render_page(prose[page_no], page_tables)
            doc.tables.extend(page_tables)
    doc.pages = [DocumentPage(source=doc.source, page=n, text=texts[n]) for n in sorted(texts)]
    doc.meta["prose"] = prose


def _page_break_before(paragraph: Paragraph) -> bool:
    p_pr = paragraph._p.pPr
    return p_pr is not None and p_pr.find(qn("w:pageBreakBefore")) is not None


def _split_paragraph(paragraph: Paragraph) -> list[str]:
    """Split paragraph text on explicit or last-rendered page breaks."""
    parts = [""]
    for el in paragraph._p.iter():
        if el.tag == qn("w:lastRenderedPageBreak"):
            parts.append("")
        elif el.tag == qn("w:br") and el.get(qn("w:type")) == "page":
            parts.append("")
        elif el.tag == qn("w:tab"):
            parts[-1] += "\t"
        elif el.tag == qn("w:t") and el.text:
            parts[-1] += el.text
    return parts


def _table_text(table: Table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [_clean(cell.text) for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def load_docx(path: Path) -> LoadedDocument:
    document = Document(str(path))
    buf: list[str] = []
    pages: list[DocumentPage] = []

    def flush() -> None:
        text = _clean("\n".join(buf))
        buf.clear()
        if text:
            pages.append(
                DocumentPage(source=str(path), page=len(pages) + 1, text=text)
            )

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            if _page_break_before(paragraph) and buf:
                flush()
            parts = _split_paragraph(paragraph)
            for i, part in enumerate(parts):
                if i > 0:
                    flush()
                cleaned = part.strip()
                if cleaned:
                    buf.append(cleaned)
        elif child.tag == qn("w:tbl"):
            rendered = _table_text(Table(child, document))
            if rendered:
                buf.append(rendered)
    flush()
    if not pages:
        raise ValueError(f"No extractable text in Word document: {path}")
    return LoadedDocument(source=str(path), pages=pages)


# --- HTML ---------------------------------------------------------------------

_HTML_NOISE_TAGS = ("script", "style", "noscript", "iframe", "nav", "header", "footer", "select", "button")
_HTML_NOISE_ATTR = re.compile(
    r"(?:^|[\s_-])(?:nav|navbar|menu|footer|foot|header|breadcrumb|banner|sidebar|copyright|share|topbar|toolbar)(?:[\s_-]|$)",
    re.I,
)
_CHARSET_RE = re.compile(rb"<meta[^>]+charset=[\"']?\s*([\w-]+)", re.I)
_CHARSET_ALIASES = {"gb2312": "gb18030", "gbk": "gb18030", "utf8": "utf-8"}


def decode_html(raw: bytes) -> str:
    """Charset from <meta>, then utf-8, then gb18030 (superset of gbk / gb2312)."""
    m = _CHARSET_RE.search(raw[:4096])
    declared = m.group(1).decode("ascii", "ignore").lower() if m else ""
    for enc in [_CHARSET_ALIASES.get(declared, declared), "utf-8", "gb18030"]:
        if not enc:
            continue
        try:
            return raw.decode(enc)
        except (LookupError, UnicodeDecodeError):
            continue
    return raw.decode("utf-8", errors="replace")


def _span(value: Any) -> int:
    m = re.match(r"\s*(\d+)", str(value or ""))
    return min(max(int(m.group(1)), 1), 1000) if m else 1


def decode_data_uri(src: str) -> tuple[str, bytes] | None:
    m = re.match(r"data:([\w/+.-]*)(;[^,]*)?,(.*)", src, re.S)
    if not m:
        return None
    mime, params, payload = m.group(1) or "", m.group(2) or "", m.group(3)
    try:
        data = base64.b64decode(re.sub(r"\s+", "", payload)) if ";base64" in params else payload.encode()
    except (binascii.Error, ValueError):
        return None
    return mime, data


def parse_html(html: str, source: str = "") -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "lxml")
    title = tbl.clean_cell(soup.title.get_text()) if soup.title else ""

    noise = list(soup.find_all(_HTML_NOISE_TAGS))
    noise += [
        el
        for el in soup.find_all(True)
        if _HTML_NOISE_ATTR.search(" ".join(el.get("class") or []) + " " + str(el.get("id") or ""))
    ]
    for el in noise:
        if not getattr(el, "decomposed", False) and el.find("table") is None and el.name not in ("html", "body"):
            el.decompose()

    images: list[ImageRef] = []
    for img in soup.find_all("img"):
        src = str(img.get("src") or "").strip()
        alt = tbl.clean_cell(img.get("alt"))
        if src.startswith("data:"):
            decoded = decode_data_uri(src)
            if decoded is None:
                continue
            mime, data = decoded
            images.append(
                ImageRef(source=source, index=len(images), mime=mime, sha256=hashlib.sha256(data).hexdigest(),
                         alt=alt, data=data, needs_ocr=mime.startswith("image/"))
            )
        elif src:
            images.append(ImageRef(source=source, index=len(images), url=src, alt=alt))

    tables: list[tbl.Table] = []
    for t in [t for t in soup.find_all("table") if t.find("table") is None]:
        rows = [
            [(tbl.clean_cell(td.get_text(" ", strip=True)), _span(td.get("rowspan")), _span(td.get("colspan")))
             for td in tr.find_all(["td", "th"], recursive=False)]
            for tr in t.find_all("tr")
        ]
        caption = t.find("caption")
        idx = len(tables)
        if any(cell for row in rows for cell, _, _ in row):
            cells, origins = tbl.expand_spans(rows)
            tables.append(
                tbl.Table(cells, source=source, page=1, index=idx,
                          title=tbl.clean_cell(caption.get_text()) if caption else "", origins=origins)
            )
            t.replace_with(soup.new_string(f"\n{tbl.TABLE_MARKER.format(index=idx)}\n"))
        else:
            t.decompose()

    root = soup.body or soup
    prose = _clean(root.get_text("\n"))
    if title and not prose.startswith(title):
        prose = f"{title}\n{prose}"
    text = tbl.render_page(prose, tables)
    return ParsedDocument(
        source=source,
        pages=[DocumentPage(source=source, page=1, text=text)] if text else [],
        tables=tables,
        images=images,
        meta={"format": "html", "title": title, "prose": {1: prose}, "ocr_pages": []},
    )


def load_html(path: Path) -> ParsedDocument:
    return parse_html(decode_html(path.read_bytes()), source=str(path))


# --- spreadsheets -------------------------------------------------------------


def _sheet_page(
    source: str,
    page_no: int,
    sheet_name: str,
    values: list[list[str]],
    merges: list[tuple[int, int, int, int]],
    first_index: int,
) -> tuple[str, list[tbl.Table]]:
    cells, origins = tbl.expand_merged_ranges(values, merges)
    prose_lines = [f"[{sheet_name}]"] if sheet_name else []
    tables: list[tbl.Table] = []
    for lines, rows, rows_o in tbl.split_blocks(cells, origins):
        title = ""
        if rows:
            width = max((c + 1 for row in rows for c, v in enumerate(row) if v), default=0)
            rows = [row[:width] for row in rows]
            rows_o = [row[:width] for row in rows_o]
            if lines:
                title = lines[-1]
                lines = lines[:-1]
        prose_lines.extend(lines)
        if rows:
            t = tbl.Table(rows, source=source, page=page_no, index=first_index + len(tables), title=title, origins=rows_o)
            tables.append(t)
            prose_lines.append(tbl.TABLE_MARKER.format(index=t.index))
    return "\n".join(prose_lines), tables


def _spreadsheet_doc(source: str, fmt: str, sheets: list[tuple[str, list[list[str]], list[tuple[int, int, int, int]]]]) -> ParsedDocument:
    pages: list[DocumentPage] = []
    tables: list[tbl.Table] = []
    prose: dict[int, str] = {}
    page_no = 0
    for name, values, merges in sheets:
        if not any(v for row in values for v in row):
            continue
        page_no += 1
        text_with_markers, sheet_tables = _sheet_page(source, page_no, name, values, merges, len(tables))
        prose[page_no] = text_with_markers
        tables.extend(sheet_tables)
        text = tbl.render_page(text_with_markers, sheet_tables)
        if text:
            pages.append(DocumentPage(source=source, page=page_no, text=text))
    if not pages:
        raise ValueError(f"No data in spreadsheet: {source}")
    return ParsedDocument(source=source, pages=pages, tables=tables, meta={"format": fmt, "prose": prose, "ocr_pages": []})


def load_xlsx(path: Path) -> ParsedDocument:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    sheets = []
    for ws in wb.worksheets:
        values = [[tbl.clean_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
        merges = [(r.min_row - 1, r.max_row, r.min_col - 1, r.max_col) for r in ws.merged_cells.ranges]
        sheets.append((ws.title, values, merges))
    return _spreadsheet_doc(str(path), "xlsx", sheets)


def load_xls(path: Path) -> ParsedDocument:
    import xlrd

    book = xlrd.open_workbook(str(path), formatting_info=True)
    sheets = []
    for sh in book.sheets():
        values = [[tbl.clean_cell(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]
        sheets.append((sh.name, values, list(sh.merged_cells)))
    return _spreadsheet_doc(str(path), "xls", sheets)


# --- images -------------------------------------------------------------------

_IMAGE_MIME = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png"}


def ocr_page(
    image: bytes, backend: OCRBackend, source: str, page_no: int, first_index: int
) -> tuple[str, list[tbl.Table]] | None:
    """OCR one image → (prose with ``[[TABLE:n]]`` markers, tables). Vision backends return
    Markdown tables; box backends get the table rebuilt from the ruling lines."""
    from doc_agent.ingest.ocr import group_lines
    from doc_agent.ingest.ocr_table import inside_bbox, rebuild_table

    result = backend.recognize(image)
    if result is None:
        return None
    tables: list[tbl.Table] = []
    lines = list(result.lines)
    if result.tables:
        for rows in result.tables:
            cells = [[tbl.clean_cell(c) for c in row] for row in rows]
            tables.append(tbl.Table(cells, source=source, page=page_no, index=first_index + len(tables)))
    elif result.boxes and (grid := rebuild_table(image, result.boxes)) is not None:
        tables.append(tbl.Table(grid.cells, source=source, page=page_no, index=first_index, origins=grid.origins))
        lines = group_lines([b for b in result.boxes if not inside_bbox(b, grid.bbox)])
    prose = "\n".join([*lines, *(tbl.TABLE_MARKER.format(index=t.index) for t in tables)])
    return prose, tables


def load_image(path: Path, *, ocr: OCRBackend | None = None) -> ParsedDocument:
    backend = ocr or NoOCR()
    data = path.read_bytes()
    image = ImageRef(
        source=str(path), index=0, mime=_IMAGE_MIME.get(path.suffix.lower(), ""),
        sha256=hashlib.sha256(data).hexdigest(), needs_ocr=True,
    )
    pages: list[DocumentPage] = []
    tables: list[tbl.Table] = []
    meta: dict[str, Any] = {"format": "image", "ocr_pages": [], "ocr_backend": backend.name, "prose": {}}
    if backend.available():
        try:
            out = ocr_page(data, backend, str(path), 1, 0)
        except Exception as exc:  # noqa: BLE001
            meta["ocr_error"] = f"{type(exc).__name__}: {exc}"
            out = None
        if out is not None:
            prose, tables = out
            text = tbl.render_page(prose, tables)
            if text:
                pages.append(DocumentPage(source=str(path), page=1, text=text))
                meta["prose"] = {1: prose}
                meta["ocr_applied"] = [1]
                image.needs_ocr = False
    return ParsedDocument(source=str(path), pages=pages, tables=tables, images=[image], meta=meta)


def render_pdf_pages(path: Path, page_numbers: list[int], scale: float = 2.0) -> dict[int, bytes]:
    """PNG bytes of the given 1-based pages (for OCR of scanned pages)."""
    import io

    import pypdfium2 as pdfium

    out: dict[int, bytes] = {}
    pdf = pdfium.PdfDocument(str(path))
    try:
        for n in page_numbers:
            buf = io.BytesIO()
            pdf[n - 1].render(scale=scale).to_pil().save(buf, format="PNG")
            out[n] = buf.getvalue()
    finally:
        pdf.close()
    return out


def _ocr_pdf_pages(path: Path, doc: ParsedDocument, backend: OCRBackend) -> None:
    """Replace scanned pages (``meta["ocr_pages"]``) with OCR text and tables, in page order."""
    done: list[int] = []
    prose_map = doc.meta.setdefault("prose", {})
    for page_no, image in render_pdf_pages(path, doc.ocr_pages).items():
        try:
            out = ocr_page(image, backend, str(path), page_no, len(doc.tables))
        except Exception as exc:  # noqa: BLE001
            doc.meta.setdefault("ocr_errors", {})[page_no] = f"{type(exc).__name__}: {exc}"
            continue
        if out is None:
            continue
        prose, page_tables = out
        text = tbl.render_page(prose, page_tables)
        if not text:
            continue
        doc.tables.extend(page_tables)
        doc.pages.append(DocumentPage(source=str(path), page=page_no, text=text))
        prose_map[page_no] = prose
        done.append(page_no)
    doc.pages.sort(key=lambda p: p.page)
    doc.meta["ocr_pages"] = [p for p in doc.ocr_pages if p not in done]
    doc.meta["ocr_applied"] = done
    doc.meta["ocr_backend"] = backend.name


def load_file(path: Path, *, tables: bool = False, ocr: OCRBackend | None = None) -> ParsedDocument:
    """Load any supported file. ``tables`` only matters for PDFs (HTML / spreadsheets
    always carry tables); ``ocr`` defaults to the no-op backend."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        doc = load_pdf(path, tables=tables)
        if ocr is not None and doc.ocr_pages and ocr.available():
            _ocr_pdf_pages(path, doc, ocr)
        return doc
    if suffix == ".docx":
        return as_parsed(load_docx(path), format="docx")
    if suffix == ".doc":
        raise ValueError(f"Legacy .doc is not supported, save as .docx: {path.name}")
    if suffix in {".txt", ".md"}:
        return as_parsed(load_txt(path), format=suffix.lstrip("."))
    if suffix in {".html", ".htm"}:
        return load_html(path)
    if suffix == ".xlsx":
        return load_xlsx(path)
    if suffix == ".xls":
        return load_xls(path)
    if suffix in _IMAGE_MIME:
        return load_image(path, ocr=ocr)
    raise ValueError(f"Unsupported file type: {path.suffix}")


_SOURCE_SUFFIXES = {
    ".pdf", ".txt", ".md", ".docx", ".doc",
    ".html", ".htm", ".xlsx", ".xls", ".jpg", ".jpeg", ".png",
}


def iter_source_files(root: Path) -> list[Path]:
    if root.is_file():
        return [root]
    files: list[Path] = []
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.suffix.lower() in _SOURCE_SUFFIXES:
            files.append(path)
    # de-dupe while preserving order
    seen: set[Path] = set()
    out: list[Path] = []
    for f in files:
        rp = f.resolve()
        if rp in seen:
            continue
        seen.add(rp)
        out.append(f)
    return out
